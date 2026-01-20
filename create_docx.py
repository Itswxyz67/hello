import os
from docx import Document
from docx.shared import Inches

def create_docx():
    doc = Document()
    doc.add_heading('Web Pages Collection', 0)

    html_dir = "html_pages"
    screenshot_dir = "screenshots"

    files = sorted([f for f in os.listdir(html_dir) if f.endswith(".html")])

    for filename in files:
        title = filename.replace(".html", "").replace("_", " ").title()
        doc.add_heading(title, level=1)

        # Add HTML Code
        doc.add_heading('HTML Code', level=2)
        with open(os.path.join(html_dir, filename), 'r') as f:
            code = f.read()
        doc.add_paragraph(code)

        # Add Screenshot
        doc.add_heading('Screenshot', level=2)
        screenshot_filename = filename.replace(".html", ".png")
        screenshot_path = os.path.join(screenshot_dir, screenshot_filename)
        if os.path.exists(screenshot_path):
            doc.add_picture(screenshot_path, width=Inches(6))
        else:
            doc.add_paragraph(f"Screenshot not found for {filename}")

        doc.add_page_break()

    doc.save('web_pages_report.docx')
    print("Document created: web_pages_report.docx")

if __name__ == "__main__":
    create_docx()
