import os
from docx import Document
from docx.shared import Inches

def generate_python_docx():
    doc = Document()

    # Add Header
    header = doc.sections[0].header
    header_para = header.paragraphs[0]
    header_para.text = "Prof. Sufiya Ahmed\nR.D. National College"

    doc.add_heading('DAA Practicals Report', 0)

    py_dir = "python_practicals"
    screenshot_dir = "python_screenshots"

    files = sorted([f for f in os.listdir(py_dir) if f.endswith(".py")])

    for filename in files:
        name = filename[:-3]
        title = name.replace("_", " ")
        doc.add_heading(title, level=1)

        # Add Source Code
        doc.add_heading('Source Code', level=2)
        with open(os.path.join(py_dir, filename), "r") as f:
            code = f.read()
        doc.add_paragraph(code)

        # Add Screenshot
        doc.add_heading('Output Screenshot', level=2)
        screenshot_path = os.path.join(screenshot_dir, name + ".png")
        if os.path.exists(screenshot_path):
            doc.add_picture(screenshot_path, width=Inches(6))
        else:
            doc.add_paragraph("Screenshot not found.")

        doc.add_page_break()

    doc.save("daa_practicals_report.docx")
    print("Created daa_practicals_report.docx")

if __name__ == "__main__":
    generate_python_docx()
