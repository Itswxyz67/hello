import os
from docx import Document
from docx.shared import Inches

def generate_docx():
    doc = Document()

    # Add Header
    header = doc.sections[0].header
    header_para = header.paragraphs[0]
    header_para.text = "Name: [Your Name]\t\tRoll No: [Your Roll No]"

    doc.add_heading('C++ Practicals Report', 0)

    cpp_dir = "cpp_practicals"
    screenshot_dir = "cpp_screenshots"

    files = sorted([f for f in os.listdir(cpp_dir) if f.endswith(".cpp")])

    for filename in files:
        name = filename[:-4]
        title = name.replace("_", " ")
        doc.add_heading(title, level=1)

        # Add Source Code
        doc.add_heading('Source Code', level=2)
        with open(os.path.join(cpp_dir, filename), "r") as f:
            code = f.read()
        doc.add_paragraph(code)

        # Add Screenshot
        doc.add_heading('Output Screenshot', level=2)
        screenshot_path = os.path.join(screenshot_dir, name + ".png")
        if os.path.exists(screenshot_path):
            doc.add_picture(screenshot_path, width=Inches(6))
        else:
            doc.add_paragraph("Screenshot not found.")

        doc.add_page_break()

    doc.save("cpp_practicals_report.docx")
    print("Created cpp_practicals_report.docx")

if __name__ == "__main__":
    generate_docx()
